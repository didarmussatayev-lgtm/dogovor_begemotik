from __future__ import annotations

import base64
import logging
import re
import subprocess
from datetime import date, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from jinja2 import TemplateSyntaxError

logger = logging.getLogger(__name__)

_JINJA_PRINT_RE = re.compile(r"\{\{.*?\}\}", flags=re.DOTALL)
_XML_TAG_RE = re.compile(r"<[^>]+>")
_SAFE_JINJA_KEY_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_PLACEHOLDER_ALIASES = {
    "дата рождения": "birth_date",
    "пол": "gender",
}


def _run_libreoffice(command: list[str], timeout: int, error_message: str) -> subprocess.CompletedProcess[str]:
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=timeout,
        )
    except FileNotFoundError as exc:
        raise RuntimeError(
            "LibreOffice is not installed or not found in PATH"
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(error_message) from exc

    if result.returncode != 0:
        logger.error("LibreOffice stderr: %s", result.stderr)
        raise RuntimeError(f"LibreOffice command failed: {result.stderr.strip()}")

    return result


def _decode_signature(signature_base64: str) -> bytes:
    """Strip optional data-URL prefix and return raw PNG bytes."""
    if "," in signature_base64:
        signature_base64 = signature_base64.split(",", 1)[1]
    return base64.b64decode(signature_base64)


def _normalize_template_expression(raw_expression: str) -> tuple[str, str | None]:
    plain_expression = _XML_TAG_RE.sub("", raw_expression)
    if not plain_expression.startswith("{{") or not plain_expression.endswith("}}"):
        return raw_expression, None

    placeholder = re.sub(r"\s+", " ", plain_expression[2:-2]).strip()
    alias_key = placeholder.lower()
    mapped = _PLACEHOLDER_ALIASES.get(alias_key)

    if mapped and mapped != placeholder:
        return f"{{{{ {mapped} }}}}", placeholder

    if _SAFE_JINJA_KEY_RE.fullmatch(placeholder):
        return raw_expression, None

    return raw_expression, placeholder


def _prepare_template_for_render(template_path: str | Path, output_dir: str | Path, output_basename: str) -> tuple[Path, list[str]]:
    source = _resolve_renderable_template(template_path=template_path, output_dir=output_dir)
    normalized_path = Path(output_dir) / f"{output_basename}_template.docx"
    rewritten_placeholders: list[str] = []
    suspicious_placeholders: list[str] = []

    with ZipFile(source, "r") as src, ZipFile(normalized_path, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                xml = data.decode("utf-8")

                def _replace(match: re.Match[str]) -> str:
                    rewritten, suspicious = _normalize_template_expression(match.group(0))
                    if rewritten != match.group(0):
                        rewritten_placeholders.append(suspicious or "")
                    elif suspicious:
                        suspicious_placeholders.append(suspicious)
                    return rewritten

                xml = _JINJA_PRINT_RE.sub(_replace, xml)
                data = xml.encode("utf-8")
            dst.writestr(info, data)

    if rewritten_placeholders:
        cleaned = sorted({p for p in rewritten_placeholders if p})
        logger.info("Normalized legacy placeholders in %s: %s", source.name, ", ".join(cleaned))

    return normalized_path, sorted(set(suspicious_placeholders))


def _resolve_renderable_template(template_path: str | Path, output_dir: str | Path) -> Path:
    source = Path(template_path)
    suffix = source.suffix.lower()
    if suffix == ".docx":
        return source
    if suffix != ".doc":
        raise RuntimeError(f"Unsupported template format: {source.suffix or '<none>'}")

    converted_path = Path(output_dir) / f"{source.stem}.docx"
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "docx",
        "--outdir",
        str(Path(output_dir)),
        str(source),
    ]
    logger.info("Converting DOC template to DOCX: %s", source)
    _run_libreoffice(
        command=cmd,
        timeout=120,
        error_message="LibreOffice DOCX conversion timed out",
    )

    if not converted_path.exists():
        raise RuntimeError(f"DOCX template not found after conversion: {converted_path}")

    logger.info("DOC template converted to DOCX: %s", converted_path)
    return converted_path


def generate_docx(
    template_path: str | Path,
    full_name: str,
    phone: str,
    iin: str,
    allergy: str,
    procedure: str,
    signature_base64: str,
    degree_of_kinship: str,
    guardian_relationship: str,
    name_surname_of_child: str,
    name_surname_patient: str,
    date_of_birth: date | None,
    id_number: str,
    id_authority: str,
    id_date_of_issue: date | None,
    adress: str,
    degree_of_kinship_mother_father_guardin: str,
    contact_name_surname_1: str,
    contact_phones_1: str,
    contact_name_surname_2: str,
    contact_phones_2: str,
    contact_name_surname_3: str,
    contact_phones_3: str,
    agreement_id: str,
    output_basename: str,
    output_dir: str | Path,
    representative_full_name: str = "",
) -> Path:
    """Fill the DOCX template and return the path to the generated file."""
    normalized_template_path, suspicious_placeholders = _prepare_template_for_render(
        template_path=template_path,
        output_dir=output_dir,
        output_basename=output_basename,
    )
    tpl = DocxTemplate(normalized_template_path)

    # Decode signature and write to a temp PNG so InlineImage can read it
    sig_bytes = _decode_signature(signature_base64)
    sig_tmp = Path(output_dir) / f"{agreement_id}_sig.png"
    sig_tmp.write_bytes(sig_bytes)

    date_of_birth_text = date_of_birth.strftime("%d.%m.%Y") if date_of_birth else ""
    id_date_of_issue_text = id_date_of_issue.strftime("%d.%m.%Y") if id_date_of_issue else ""
    now = datetime.now()

    context = {
        "full_name": full_name,
        "name_surname": full_name,
        "phone": phone,
        "iin": iin,
        "allergy": allergy,
        "procedure": procedure,
        "degree_of_kinship": degree_of_kinship,
        "guardian_relationship": guardian_relationship,
        "name_surname_of_child": name_surname_of_child,
        "name_surname_patient": name_surname_patient,
        "date_of_birth": date_of_birth_text,
        "birth_date": date_of_birth_text,
        "gender": "",
        "id_number": id_number,
        "id_authority": id_authority,
        "id_date_of_issue": id_date_of_issue_text,
        "adress": adress,
        "degree_of_kinship_mother_father_guardin": degree_of_kinship_mother_father_guardin,
        "contact_name_surname_1": contact_name_surname_1,
        "contact_phones_1": contact_phones_1,
        "contact_name_surname_2": contact_name_surname_2,
        "contact_phones_2": contact_phones_2,
        "contact_name_surname_3": contact_name_surname_3,
        "contact_phones_3": contact_phones_3,
        "contact_name_surname": contact_name_surname_1,
        "contact_phones": contact_phones_1,
        "representative_full_name": representative_full_name,
        "date": date.today().strftime("%d.%m.%Y"),
        "full_date": now.strftime("%d.%m.%Y %H часов %M минут"),
        "agreement_id": agreement_id,
        "signature": InlineImage(tpl, str(sig_tmp), width=Mm(50)),
    }

    try:
        tpl.render(context)
    except TemplateSyntaxError as exc:
        hint = ""
        if suspicious_placeholders:
            hint = (
                f" Likely invalid placeholder(s): {', '.join(suspicious_placeholders[:3])}. "
                "Use ASCII underscore keys like {{ birth_date }} or {{ gender }}."
            )
        raise RuntimeError(f"Template syntax error in {Path(template_path).name}: {exc}.{hint}") from exc

    docx_path = Path(output_dir) / f"{output_basename}.docx"
    tpl.save(str(docx_path))
    logger.info("DOCX generated: %s", docx_path)
    return docx_path


def generate_begemotik_docx(
    template_path: str | Path,
    iin: str,
    surname: str,
    name: str,
    last_name: str,
    gender: str,
    birthdate: str,
    phone: str,
    has_kinship: bool,
    surname_kinship: str,
    name_kinship: str,
    last_name_kinship: str,
    degree_of_kinship: str,
    allergy_value: str,
    no_allergy_value: str,
    procedure: str,
    signature_base64: str,
    agreement_id: str,
    output_basename: str,
    output_dir: str | Path,
) -> Path:
    """Fill begemotik_template.docx and return path to generated DOCX."""
    normalized_template_path, suspicious_placeholders = _prepare_template_for_render(
        template_path=template_path,
        output_dir=output_dir,
        output_basename=output_basename,
    )
    tpl = DocxTemplate(normalized_template_path)

    sig_bytes = _decode_signature(signature_base64)
    sig_tmp = Path(output_dir) / f"{agreement_id}_sig.png"
    sig_tmp.write_bytes(sig_bytes)

    now = datetime.now()

    # Build full names
    patient_full_name = " ".join(filter(None, [surname, name, last_name]))
    rep_full_name = " ".join(filter(None, [surname_kinship, name_kinship, last_name_kinship])) if has_kinship else ""

    # name_surname_kinship: FIO of the patient being represented (used in consent sentence)
    name_surname_kinship = patient_full_name if has_kinship else ""

    # patient / kinship fields (only one filled at a time)
    if has_kinship:
        patient_field = ""
        kinship_field = rep_full_name
    else:
        patient_field = patient_full_name
        kinship_field = ""

    # signature_kinship: representative's FIO when present, else empty
    signature_kinship = rep_full_name if has_kinship else ""

    context = {
        "iin": iin,
        "surname": surname,
        "name": name,
        "last_name": last_name,
        "gender": gender,
        "birthdate": birthdate,
        "phone": phone,
        # kinship
        "surname_kinship": surname_kinship,
        "name_kinship": name_kinship,
        "last_name_kinship": last_name_kinship,
        "degree_of_kinship": degree_of_kinship,
        "name_surname_kinship": name_surname_kinship,
        "signature_kinship": signature_kinship,
        # patient / kinship signer disambiguation
        "patient": patient_field,
        "kinship": kinship_field,
        # allergy
        "allergy": allergy_value,
        "no_allergy": no_allergy_value,
        # procedure
        "procedure": procedure,
        # date/time
        "date": now.strftime("%d.%m.%Y"),
        "full_date": now.strftime("%d.%m.%Y %H:%M"),
        "agreement_id": agreement_id,
        # signature image
        "signature": InlineImage(tpl, str(sig_tmp), width=Mm(50)),
    }

    try:
        tpl.render(context)
    except TemplateSyntaxError as exc:
        hint = ""
        if suspicious_placeholders:
            hint = (
                f" Likely invalid placeholder(s): {', '.join(suspicious_placeholders[:3])}. "
                "Use ASCII underscore keys like {{ birth_date }} or {{ gender }}."
            )
        raise RuntimeError(f"Template syntax error in {Path(template_path).name}: {exc}.{hint}") from exc

    docx_path = Path(output_dir) / f"{output_basename}.docx"
    tpl.save(str(docx_path))
    logger.info("Begemotik DOCX generated: %s", docx_path)
    return docx_path


def convert_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    """Convert a DOCX file to PDF using LibreOffice headless."""
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(docx_path),
    ]
    logger.info("Running LibreOffice: %s", " ".join(cmd))
    _run_libreoffice(
        command=cmd,
        timeout=120,
        error_message="LibreOffice conversion timed out",
    )

    pdf_path = output_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError(f"PDF not found after conversion: {pdf_path}")

    logger.info("PDF generated: %s", pdf_path)
    return pdf_path
